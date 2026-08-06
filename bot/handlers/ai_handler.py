"""AI mutaxassis: suhbat va hisobotlarni umumlashtirish."""
from __future__ import annotations

import io
import logging
from datetime import datetime

from aiogram import Bot, F, Router
from aiogram.filters import Command, CommandObject
from aiogram.types import (
    BufferedInputFile,
    CallbackQuery,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    Message,
)

from ..ai import AIClient
from ..config import Config
from ..database import Database
from ..utils.extract import extract_text

logger = logging.getLogger(__name__)

router = Router()


def _ai_or_none(config: Config, ai: AIClient | None) -> AIClient | None:
    if not config.ai_enabled or ai is None:
        return None
    return ai


# ── /ai — erkin suhbat ───────────────────────────────────────

@router.message(Command("ai"))
async def cmd_ai_chat(
    message: Message, command: CommandObject, config: Config, db: Database,
    ai: AIClient | None = None,
) -> None:
    client = _ai_or_none(config, ai)
    if not client:
        await message.reply(
            "🤖 AI funksiyalari o'chirilgan.\n"
            "<code>ANTHROPIC_API_KEY</code> ni .env faylida to'ldiring."
        )
        return
    question = (command.args or "").strip()
    if not question:
        await message.reply(
            "ℹ️ Foydalanish: <code>/ai savol matni</code>\n\n"
            "Masalan: <code>/ai Oylik hisobot topshirig'ini qanday yozaman?</code>"
        )
        return

    wait = await message.reply("🤖 O'ylamoqda...")

    extra = await _build_context(db, config, message.from_user.id if message.from_user else 0)
    messages = [{"role": "user", "content": question}]
    answer = await client.chat(messages, extra_context=extra)

    try:
        await wait.edit_text(f"🤖 <b>AI javob:</b>\n\n{answer}")
    except Exception:
        await message.reply(f"🤖 <b>AI javob:</b>\n\n{answer}")


async def _build_context(db: Database, config: Config, user_id: int) -> str:
    parts = []
    tasks = await db.list_open_tasks()
    if tasks:
        lines = []
        for t in tasks:
            submitted = await db.submitted_employee_ids(t["id"])
            total = await db.count_employees()
            lines.append(f"#{t['id']} {t['title']} — {len(submitted)}/{total} bajardi")
        parts.append("Ochiq topshiriqlar:\n" + "\n".join(lines))
    emp = await db.get_employee(user_id)
    if emp:
        parts.append(f"Foydalanuvchi: {emp['full_name']} (xodim)")
    elif config.is_manager(user_id):
        parts.append("Foydalanuvchi: Rahbar")
    return "\n\n".join(parts)


# ── /umumlashtir — hisobotlarni umumlashtirish ───────────────

@router.message(Command("umumlashtir", "consolidate"))
async def cmd_consolidate(
    message: Message, command: CommandObject, config: Config, db: Database,
    bot: Bot, ai: AIClient | None = None,
) -> None:
    if not (message.from_user and config.is_manager(message.from_user.id)):
        await message.reply("⛔️ Bu komanda faqat rahbar uchun.")
        return
    client = _ai_or_none(config, ai)
    if not client:
        await message.reply(
            "🤖 AI funksiyalari o'chirilgan.\n"
            "<code>ANTHROPIC_API_KEY</code> ni .env faylida to'ldiring."
        )
        return

    if not (command.args and command.args.strip().isdigit()):
        await message.reply(
            "ℹ️ Foydalanish: <code>/umumlashtir N</code>\n"
            "N — topshiriq raqami. Xodimlar yuborgan fayllarni AI umumlashtiradi."
        )
        return

    task_id = int(command.args.strip())
    await _do_consolidate(message, bot, db, config, client, task_id)


@router.callback_query(F.data.startswith("consolidate:"))
async def cb_consolidate(
    callback: CallbackQuery, db: Database, config: Config, bot: Bot,
    ai: AIClient | None = None,
) -> None:
    if not config.is_manager(callback.from_user.id):
        await callback.answer("⛔️ Faqat rahbar uchun.", show_alert=True)
        return
    client = _ai_or_none(config, ai)
    if not client:
        await callback.answer("AI funksiyalari o'chirilgan.", show_alert=True)
        return
    task_id = int(callback.data.split(":")[1])
    await callback.answer("🤖 Umumlashtirish boshlanmoqda...")
    await _do_consolidate(callback.message, bot, db, config, client, task_id)


async def _do_consolidate(
    reply_to: Message, bot: Bot, db: Database, config: Config,
    client: AIClient, task_id: int,
) -> None:
    task = await db.get_task(task_id)
    if not task:
        await reply_to.reply("⚠️ Bunday topshiriq topilmadi.")
        return

    submissions = await db.get_submissions(task_id)
    if not submissions:
        await reply_to.reply(
            f"⚠️ #{task_id}-topshiriq bo'yicha hali hech kim ish topshirmagan."
        )
        return

    wait = await reply_to.reply(
        f"🤖 #{task_id}-topshiriq bo'yicha {len(submissions)} ta hisobotni "
        f"yuklab umumlashtirmoqda..."
    )

    reports: list[tuple[str, str]] = []
    for sub in submissions:
        emp = await db.get_employee(sub["employee_id"])
        emp_name = emp["full_name"] if emp else f"ID:{sub['employee_id']}"

        text_parts = []
        if sub["note"]:
            text_parts.append(sub["note"])

        if sub["file_id"] and sub["file_name"]:
            try:
                tg_file = await bot.get_file(sub["file_id"])
                data = io.BytesIO()
                await bot.download_file(tg_file.file_path, data)
                extracted = extract_text(data.getvalue(), sub["file_name"])
                text_parts.append(f"[Fayl: {sub['file_name']}]\n{extracted}")
            except Exception as exc:
                logger.warning("Fayl yuklashda xato (sub %s): %s", sub["id"], exc)
                text_parts.append(f"[Faylni yuklashda xatolik: {exc}]")

        combined = "\n\n".join(text_parts) if text_parts else "[Bo'sh hisobot]"
        reports.append((emp_name, combined))

    if not reports:
        try:
            await wait.edit_text("⚠️ Umumlashtiradigan hisobot topilmadi.")
        except Exception:
            pass
        return

    result = await client.consolidate(
        task_title=task["title"],
        task_desc=task["description"] or "",
        reports=reports,
    )

    if not result:
        try:
            await wait.edit_text(
                "⚠️ AI umumlashtirishda xatolik yuz berdi. Keyinroq urinib ko'ring."
            )
        except Exception:
            pass
        return

    summary_text = (
        f"🤖 <b>AI UMUMLASHTIRISH — #{task_id}: {task['title']}</b>\n\n"
        f"📝 <b>Xulosa:</b>\n{result.get('summary', '—')}\n\n"
    )
    key_points = result.get("key_points", [])
    if key_points:
        summary_text += "🔑 <b>Asosiy natijalar:</b>\n"
        for kp in key_points:
            summary_text += f"  • {kp}\n"
        summary_text += "\n"

    summary_text += f"📊 Tahlil qilingan hisobotlar: {len(reports)} ta"

    try:
        await wait.edit_text(summary_text)
    except Exception:
        await reply_to.reply(summary_text)

    columns = result.get("columns", [])
    table_rows = result.get("rows", [])
    if columns and table_rows:
        excel_data = _build_consolidation_excel(
            task, columns, table_rows, config
        )
        fname = (
            f"umumiy_{task_id}_{datetime.now(config.tz).strftime('%Y%m%d_%H%M')}.xlsx"
        )
        await bot.send_document(
            reply_to.chat.id,
            BufferedInputFile(excel_data, filename=fname),
            caption=f"📄 #{task_id} umumlashtirilgan jadval ({len(table_rows)} xodim)",
        )


def _build_consolidation_excel(task, columns, rows, config) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = f"Topshiriq #{task['id']}"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for j, col_name in enumerate(columns, 1):
        c = ws.cell(row=1, column=j, value=col_name)
        c.fill = header_fill
        c.font = header_font
        c.alignment = center

    for i, row_data in enumerate(rows, 2):
        for j, val in enumerate(row_data, 1):
            c = ws.cell(row=i, column=j, value=str(val) if val else "")
            c.alignment = Alignment(wrap_text=True, vertical="top")

    for j in range(1, len(columns) + 1):
        ws.column_dimensions[chr(64 + j) if j <= 26 else "AA"].width = 20
    ws.column_dimensions["A"].width = 25
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 30

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
