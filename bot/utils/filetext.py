"""Fayl ichidagi matnni ajratib olish (docx, xlsx, pdf) — kalit so'z matching uchun."""
from __future__ import annotations

import io

_MAX_CHARS = 3000  # Matching uchun yetarli hajm


def extract_text_from_bytes(data: bytes, file_name: str | None) -> str:
    """Fayl baytlaridan matn chiqaradi. Xato bo'lsa — bo'sh satr."""
    if not file_name:
        return ""
    ext = file_name.lower().rsplit(".", 1)[-1] if "." in file_name else ""
    try:
        if ext == "docx":
            return _read_docx(data)[:_MAX_CHARS]
        if ext in ("xlsx", "xls"):
            return _read_xlsx(data)[:_MAX_CHARS]
        if ext == "pdf":
            return _read_pdf(data)[:_MAX_CHARS]
    except Exception:
        pass
    return ""


def _read_docx(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text.strip() for p in doc.paragraphs[:40] if p.text.strip()]
    return " ".join(parts)


def _read_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in list(wb.worksheets)[:3]:
        for row in ws.iter_rows(max_row=30, values_only=True):
            for cell in row:
                if isinstance(cell, str) and cell.strip():
                    parts.append(cell.strip())
            if len(" ".join(parts)) >= _MAX_CHARS:
                break
    return " ".join(parts)


def _read_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages[:3]:
        t = page.extract_text()
        if t:
            parts.append(t)
    return " ".join(parts)
