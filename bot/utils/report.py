"""Svodka (hisobot) tayyorlash — matn va Excel formatida."""
from __future__ import annotations

import io
from datetime import datetime
from zoneinfo import ZoneInfo

import aiosqlite
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .deadline import format_deadline


def _display_name(emp: aiosqlite.Row) -> str:
    name = emp["full_name"]
    if emp["username"]:
        return f"{name} (@{emp['username']})"
    return name


def build_task_text_report(
    task: aiosqlite.Row,
    employees: list[aiosqlite.Row],
    submitted_ids: set[int],
    tz: ZoneInfo,
    partial_ids: set[int] = frozenset(),
) -> str:
    """Bitta topshiriq bo'yicha 3 guruhli svodka: to'liq / chala / bajarmagan.

    submitted_ids — to'liq bajarganlar (fully done).
    partial_ids   — chala bajarganlar (topshirishgan, lekin fayllar yetarli emas).
    """
    done    = [e for e in employees if e["tg_id"] in submitted_ids]
    partial = [e for e in employees if e["tg_id"] in partial_ids]
    all_submitted = submitted_ids | partial_ids
    not_done = [e for e in employees if e["tg_id"] not in all_submitted]
    total = len(employees)
    pct = round(len(done) / total * 100) if total else 0

    req = task["required_files"] if "required_files" in task.keys() else 0

    lines = [
        f"📋 <b>Topshiriq #{task['id']}: {task['title']}</b>",
        f"🗓 Muddat: {format_deadline(task['deadline'], tz)}",
    ]
    if req:
        lines.append(f"📎 Talab qilinadigan fayllar: {req} ta")
    lines += [
        f"📊 To'liq bajardi: <b>{len(done)}/{total}</b> ({pct}%)",
        "",
    ]
    if done:
        lines.append("✅ <b>Bajarganlar:</b>")
        for i, e in enumerate(done, 1):
            lines.append(f"  {i}. {_display_name(e)}")
        lines.append("")
    if partial:
        lines.append("⏳ <b>Chala bajarganlar:</b>")
        for i, e in enumerate(partial, 1):
            lines.append(f"  {i}. {_display_name(e)}")
        lines.append("")
    if not_done:
        lines.append("❌ <b>Bajarmaganlar:</b>")
        for i, e in enumerate(not_done, 1):
            lines.append(f"  {i}. {_display_name(e)}")
    if not employees:
        lines.append("⚠️ Xodimlar ro'yxati bo'sh. /hodim_qoshish bilan qo'shing.")
    return "\n".join(lines)


def build_overall_text_report(
    rows: list[tuple],
    employees: list[aiosqlite.Row],
    tz: ZoneInfo,
) -> str:
    """Barcha ochiq topshiriqlar bo'yicha umumiy svodka.

    rows elementlari: (task, fully_done_ids) yoki (task, fully_done_ids, partial_ids).
    """
    total_emp = len(employees)
    emp_ids = {e["tg_id"] for e in employees}
    lines = [f"📈 <b>UMUMIY SVODKA</b> — {datetime.now(tz).strftime('%d.%m.%Y %H:%M')}", ""]
    if not rows:
        lines.append("Hozircha ochiq topshiriqlar yo'q.")
        return "\n".join(lines)
    for entry in rows:
        task = entry[0]
        done_ids = entry[1]
        partial_ids: set[int] = entry[2] if len(entry) > 2 else set()
        done    = len(done_ids & emp_ids)
        partial = len(partial_ids & emp_ids)
        pct = round(done / total_emp * 100) if total_emp else 0
        bar = _progress_bar(pct)
        partial_txt = f" | ⏳{partial}" if partial else ""
        lines.append(
            f"#{task['id']} {task['title']}\n"
            f"   {bar} ✅{done}{partial_txt}/{total_emp} ({pct}%) · "
            f"{format_deadline(task['deadline'], tz)}"
        )
    return "\n".join(lines)


def _progress_bar(pct: int, width: int = 10) -> str:
    filled = round(pct / 100 * width)
    return "▓" * filled + "░" * (width - filled)


def build_excel_report(
    rows: list[tuple[aiosqlite.Row, set[int], dict[int, aiosqlite.Row]]],
    employees: list[aiosqlite.Row],
    tz: ZoneInfo,
) -> bytes:
    """Excel svodka: qatorlar — xodimlar, ustunlar — topshiriqlar.

    rows: (task, submitted_ids, {employee_id: submission_row}) ro'yxati.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Svodka"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    done_fill = PatternFill("solid", fgColor="C6EFCE")
    miss_fill = PatternFill("solid", fgColor="FFC7CE")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Sarlavha qatori
    ws.cell(row=1, column=1, value="№")
    ws.cell(row=1, column=2, value="Xodim")
    for j, (task, _subs, _detail) in enumerate(rows, start=3):
        c = ws.cell(row=1, column=j, value=f"#{task['id']}\n{task['title']}")
        c.alignment = center
    ws.cell(row=1, column=3 + len(rows), value="Jami bajardi")

    for col in range(1, 4 + len(rows)):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center

    # Xodimlar qatorlari
    for i, emp in enumerate(employees, start=1):
        r = i + 1
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=2, value=_display_name(emp))
        done_count = 0
        for j, (task, submitted, detail) in enumerate(rows, start=3):
            cell = ws.cell(row=r, column=j)
            if emp["tg_id"] in submitted:
                cell.value = "✅"
                cell.fill = done_fill
                done_count += 1
            else:
                cell.value = "—"
                cell.fill = miss_fill
            cell.alignment = center
        ws.cell(row=r, column=3 + len(rows), value=done_count).alignment = center

    # Yakuniy qator — har topshiriq bo'yicha jami
    last = len(employees) + 2
    ws.cell(row=last, column=2, value="JAMI").font = Font(bold=True)
    for j, (task, submitted, detail) in enumerate(rows, start=3):
        cnt = len(submitted & {e["tg_id"] for e in employees})
        c = ws.cell(row=last, column=j, value=f"{cnt}/{len(employees)}")
        c.font = Font(bold=True)
        c.alignment = center

    # Ustun kengliklari
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 28
    for j in range(3, 3 + len(rows)):
        ws.column_dimensions[get_column_letter(j)].width = 16
    ws.column_dimensions[get_column_letter(3 + len(rows))].width = 12
    ws.freeze_panes = "C2"
    ws.row_dimensions[1].height = 40

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
