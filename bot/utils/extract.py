"""Word/Excel/PowerPoint/PDF fayllardan matn ajratib olish (umumlashtirish uchun)."""
from __future__ import annotations

import io

MAX_CHARS = 8000  # bitta fayldan olinadigan maksimal belgi (token nazorati uchun)


def extract_text(data: bytes, filename: str | None) -> str:
    """Fayl baytlaridan matn qaytaradi. Format nomiga qarab aniqlanadi."""
    name = (filename or "").lower()
    try:
        if name.endswith(".docx"):
            text = _from_docx(data)
        elif name.endswith(".xlsx"):
            text = _from_xlsx(data)
        elif name.endswith(".pptx"):
            text = _from_pptx(data)
        elif name.endswith(".pdf"):
            text = _from_pdf(data)
        elif name.endswith((".txt", ".csv", ".md")):
            text = data.decode("utf-8", errors="ignore")
        else:
            return "[Bu fayl formatidan matn o'qib bo'lmaydi]"
    except Exception as exc:  # noqa: BLE001
        return f"[Faylni o'qishda xatolik: {exc}]"
    text = text.strip()
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS] + "\n…[qisqartirildi]"
    return text or "[Faylda matn topilmadi]"


def _from_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _from_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"# Varaq: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts)


def _from_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"# Slayd {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs).strip()
                    if line:
                        parts.append(line)
    return "\n".join(parts)


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages:
        txt = page.extract_text() or ""
        if txt.strip():
            parts.append(txt.strip())
    return "\n".join(parts)
